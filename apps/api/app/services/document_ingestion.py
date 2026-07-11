from dataclasses import dataclass
from collections.abc import Sequence
from pathlib import Path
import re

from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.config import Settings
from app.models.document import Document, DocumentChunk
from app.services.embeddings import EmbeddingError, embed_document_chunks, get_embedding_provider

DOCUMENT_STATUS_PENDING = "pending"
DOCUMENT_STATUS_PROCESSING = "processing"
DOCUMENT_STATUS_READY = "ready"
DOCUMENT_STATUS_FAILED = "failed"
TOKEN_RE = re.compile(r"\S+")
MARKDOWN_HEADING_RE = re.compile(r"^#{1,6}\s+(.+?)\s*#*\s*$")


class DocumentIngestionError(Exception):
    pass


@dataclass(frozen=True)
class TextChunk:
    content: str
    char_start: int
    char_end: int
    token_start: int
    token_end: int
    source_page: int | None
    source_label: str | None


@dataclass(frozen=True)
class SourceSpan:
    char_start: int
    char_end: int
    source_page: int | None = None
    source_label: str | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    source_spans: list[SourceSpan]


@dataclass(frozen=True)
class TokenSpan:
    text: str
    char_start: int
    char_end: int


def extract_text(document: Document) -> ExtractedDocument:
    path = Path(document.storage_path)
    if document.content_type == "text/plain":
        text = clean_extracted_text(path.read_text(encoding="utf-8"))
        return ExtractedDocument(
            text=text,
            source_spans=[SourceSpan(0, len(text), source_label=document.original_filename)],
        )

    if document.content_type == "text/markdown":
        return extract_markdown_text(path)

    if document.content_type == "application/pdf":
        return extract_pdf_text(path)

    if document.content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_docx_text(path)

    raise DocumentIngestionError(f"Unsupported content type: {document.content_type}")


def extract_markdown_text(path: Path) -> ExtractedDocument:
    raw_text = path.read_text(encoding="utf-8")
    lines: list[str] = []
    pending_labels: list[str | None] = []
    current_label: str | None = None

    for raw_line in raw_text.splitlines():
        line = clean_line(raw_line)
        heading_match = MARKDOWN_HEADING_RE.match(line)
        if heading_match:
            current_label = heading_match.group(1).strip()
        lines.append(line)
        pending_labels.append(current_label)

    return build_extracted_document(lines, pending_labels)


def extract_pdf_text(path: Path) -> ExtractedDocument:
    try:
        from pypdf import PdfReader  # type: ignore[import-not-found]
    except ImportError as exc:
        raise DocumentIngestionError("PDF extraction requires pypdf") from exc

    try:
        reader = PdfReader(str(path))
        pages = [clean_extracted_text(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:
        raise DocumentIngestionError("PDF text extraction failed") from exc

    lines: list[str] = []
    page_numbers: list[int | None] = []
    for page_number, page_text in enumerate(pages, start=1):
        if not page_text:
            continue
        if lines:
            lines.append("")
            page_numbers.append(None)
        for line in page_text.splitlines():
            lines.append(line)
            page_numbers.append(page_number)

    return build_extracted_document(lines, page_numbers, use_page_numbers=True)


def extract_docx_text(path: Path) -> ExtractedDocument:
    try:
        from docx import Document as DocxDocument  # type: ignore[import-untyped]
    except ImportError as exc:
        raise DocumentIngestionError("DOCX extraction requires python-docx") from exc

    try:
        docx_document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentIngestionError("DOCX text extraction failed") from exc

    lines: list[str] = []
    labels: list[str | None] = []
    current_label: str | None = None
    for paragraph in docx_document.paragraphs:
        line = clean_line(paragraph.text)
        if not line:
            continue
        if paragraph.style is not None and paragraph.style.name.lower().startswith("heading"):
            current_label = line
        lines.append(line)
        labels.append(current_label)

    return build_extracted_document(lines, labels)


def chunk_text(
    extracted_document: ExtractedDocument,
    chunk_size_chars: int,
    overlap_chars: int,
) -> list[TextChunk]:
    normalized_text = extracted_document.text.strip()
    if not normalized_text:
        raise DocumentIngestionError("No extractable text found")

    if chunk_size_chars <= 0:
        raise ValueError("chunk_size_chars must be greater than zero")
    if overlap_chars < 0:
        raise ValueError("overlap_chars cannot be negative")
    if overlap_chars >= chunk_size_chars:
        raise ValueError("overlap_chars must be smaller than chunk_size_chars")

    tokens = tokenize_with_spans(normalized_text)
    if not tokens:
        raise DocumentIngestionError("No extractable text found")

    max_tokens = max(1, chunk_size_chars // 4)
    overlap_tokens = min(max(0, overlap_chars // 4), max_tokens - 1)
    chunks: list[TextChunk] = []
    token_start = 0

    while token_start < len(tokens):
        token_end = min(token_start + max_tokens, len(tokens))
        char_start = tokens[token_start].char_start
        char_end = tokens[token_end - 1].char_end
        content = normalized_text[char_start:char_end].strip()
        if content:
            source_span = find_source_span(
                extracted_document.source_spans,
                char_start=char_start,
                char_end=char_end,
            )
            chunks.append(
                TextChunk(
                    content=content,
                    char_start=char_start,
                    char_end=char_end,
                    token_start=token_start,
                    token_end=token_end,
                    source_page=source_span.source_page if source_span is not None else None,
                    source_label=source_span.source_label if source_span is not None else None,
                )
            )

        if token_end >= len(tokens):
            break
        token_start = max(token_end - overlap_tokens, token_start + 1)

    if not chunks:
        raise DocumentIngestionError("No extractable text found")

    return chunks


def estimate_token_count(text: str) -> int:
    # This is a rough boundary for storage and retrieval planning until tokenizer-aware chunking lands.
    return max(1, (len(text) + 3) // 4) if text else 0


def tokenize_with_spans(text: str) -> list[TokenSpan]:
    return [TokenSpan(match.group(0), match.start(), match.end()) for match in TOKEN_RE.finditer(text)]


def clean_line(line: str) -> str:
    return re.sub(r"[ \t]+", " ", line).strip()


def clean_extracted_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [clean_line(line) for line in normalized.splitlines()]
    return collapse_blank_lines(lines)


def collapse_blank_lines(lines: list[str]) -> str:
    cleaned_lines: list[str] = []
    previous_blank = False
    for line in lines:
        is_blank = not line
        if is_blank and previous_blank:
            continue
        cleaned_lines.append(line)
        previous_blank = is_blank
    return "\n".join(cleaned_lines).strip()


def build_extracted_document(
    lines: list[str],
    metadata: Sequence[int | str | None],
    *,
    use_page_numbers: bool = False,
) -> ExtractedDocument:
    text_parts: list[str] = []
    source_spans: list[SourceSpan] = []
    cursor = 0

    for line, source_metadata in zip(lines, metadata, strict=True):
        cleaned_line = clean_line(line)
        if not cleaned_line:
            if text_parts and text_parts[-1] != "\n":
                text_parts.append("\n")
                cursor += 1
            continue

        if text_parts and text_parts[-1] != "\n":
            text_parts.append("\n")
            cursor += 1

        start = cursor
        text_parts.append(cleaned_line)
        cursor += len(cleaned_line)
        source_spans.append(
            SourceSpan(
                start,
                cursor,
                source_page=int(source_metadata) if use_page_numbers and source_metadata is not None else None,
                source_label=str(source_metadata) if not use_page_numbers and source_metadata is not None else None,
            )
        )

    text = "".join(text_parts).strip()
    return ExtractedDocument(text=text, source_spans=source_spans)


def find_source_span(
    source_spans: list[SourceSpan],
    *,
    char_start: int,
    char_end: int,
) -> SourceSpan | None:
    best_span: SourceSpan | None = None
    best_overlap = 0
    for span in source_spans:
        overlap = min(char_end, span.char_end) - max(char_start, span.char_start)
        if overlap > best_overlap:
            best_span = span
            best_overlap = overlap
    return best_span


def ingest_document(
    db: Session,
    document: Document,
    *,
    chunk_size_chars: int,
    overlap_chars: int,
) -> Document:
    document.status = DOCUMENT_STATUS_PROCESSING
    document.ingestion_error = None
    document.chunk_count = 0
    db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id))
    db.commit()
    db.refresh(document)

    try:
        extracted_document = extract_text(document)
        chunks = chunk_text(
            extracted_document,
            chunk_size_chars=chunk_size_chars,
            overlap_chars=overlap_chars,
        )
    except Exception as exc:
        document.status = DOCUMENT_STATUS_FAILED
        document.ingestion_error = str(exc)
        document.chunk_count = 0
        db.commit()
        db.refresh(document)
        return document

    for index, chunk in enumerate(chunks):
        db.add(
            DocumentChunk(
                document_id=document.id,
                chunk_index=index,
                content=chunk.content,
                char_start=chunk.char_start,
                char_end=chunk.char_end,
                token_start=chunk.token_start,
                token_end=chunk.token_end,
                source_page=chunk.source_page,
                source_label=chunk.source_label,
            )
        )

    document.status = DOCUMENT_STATUS_READY
    document.ingestion_error = None
    document.chunk_count = len(chunks)
    db.commit()
    db.refresh(document)
    return document


def process_document_ingestion(
    db: Session,
    document_id: int,
    *,
    settings: Settings,
) -> Document | None:
    document = db.get(Document, document_id)
    if document is None:
        return None
    if document.status == DOCUMENT_STATUS_READY:
        return document

    document = ingest_document(
        db,
        document,
        chunk_size_chars=settings.document_chunk_size_chars,
        overlap_chars=settings.document_chunk_overlap_chars,
    )
    if document.status != DOCUMENT_STATUS_READY:
        return document

    try:
        embed_document_chunks(db, document, get_embedding_provider())
    except EmbeddingError as exc:
        document.status = DOCUMENT_STATUS_FAILED
        document.ingestion_error = str(exc)
        db.commit()

    db.refresh(document)
    return document
