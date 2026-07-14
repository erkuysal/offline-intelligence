from pathlib import Path
from uuid import uuid4
from io import BytesIO

from docx import Document as DocxDocument
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, text

from app.config import get_settings
from app.db.session import SessionLocal
from app.main import app
from app.models.document import DocumentChunk
from app.models.retrieval import RetrievalRun
from app.services.document_ingestion import process_document_ingestion
from app.services.embeddings import EmbeddingError, FakeEmbeddingProvider, reembed_all_document_chunks

client = TestClient(app)
pytestmark = pytest.mark.usefixtures("clean_database")


def get_access_token() -> str:
    email = f"document-user-{uuid4().hex}@example.com"
    return create_user_and_get_token(email)


def create_user_and_get_token(email: str) -> str:
    password = "correct-horse-battery-staple"
    client.post(
        "/api/v1/auth/register",
        json={
            "email": email,
            "password": password,
        },
    )
    response = client.post(
        "/api/v1/auth/login",
        json={
            "email": email,
            "password": password,
        },
    )
    return response.json()["access_token"]


def upload_txt_document(token: str) -> dict:
    return upload_txt_document_with_content(token, b"Backups run every night.")


def upload_txt_document_with_content(token: str, content: bytes, filename: str = "policy.txt") -> dict:
    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                filename,
                content,
                "text/plain",
            )
        },
    )
    assert response.status_code == 201
    return response.json()


def test_list_documents_requires_authentication() -> None:
    response = client.get("/api/v1/documents")

    assert response.status_code == 401


def test_list_documents_for_authenticated_user() -> None:
    token = get_access_token()

    response = client.get(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 200
    assert response.json() == []


def test_upload_txt_document(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))

    token = get_access_token()
    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "policy.txt",
                b"Backups run every night.",
                "text/plain",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["original_filename"] == "policy.txt"
    assert body["content_type"] == "text/plain"
    assert body["size_bytes"] == len(b"Backups run every night.")
    assert body["checksum_sha256"]
    assert body["status"] == "ready"
    assert body["chunk_count"] == 1
    assert body["ingestion_error"] is None
    assert body["version_number"] == 1
    assert len(list(tmp_path.iterdir())) == 1

    list_response = client.get(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert list_response.status_code == 200
    assert len(list_response.json()) == 1

    chunks_response = client.get(
        f"/api/v1/documents/{body['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert chunks_response.status_code == 200
    chunks = chunks_response.json()
    assert len(chunks) == 1
    assert chunks[0]["chunk_index"] == 0
    assert chunks[0]["content"] == "Backups run every night."
    assert chunks[0]["char_start"] == 0
    assert chunks[0]["char_end"] == len("Backups run every night.")
    assert chunks[0]["token_start"] == 0
    assert chunks[0]["token_end"] > 0
    assert chunks[0]["embedding_model"] == "fake-bow"


def test_upload_rejects_duplicate_document_version(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    upload_txt_document(token)

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "policy.txt",
                b"Backups run every night.",
                "text/plain",
            )
        },
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "Duplicate document upload"
    assert len(list(tmp_path.iterdir())) == 1


def test_upload_new_document_version_replaces_chunks(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    document = upload_txt_document(token)

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "policy.txt",
                b"Retention reports run monthly.",
                "text/plain",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["id"] == document["id"]
    assert body["version_number"] == 2
    assert body["checksum_sha256"] != document["checksum_sha256"]

    chunks_response = client.get(
        f"/api/v1/documents/{document['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )
    versions_response = client.get(
        f"/api/v1/documents/{document['id']}/versions",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert chunks_response.status_code == 200
    assert chunks_response.json()[0]["content"] == "Retention reports run monthly."
    assert versions_response.status_code == 200
    assert [version["version_number"] for version in versions_response.json()] == [1, 2]
    assert len(list(tmp_path.iterdir())) == 2


def test_upload_txt_document_creates_multiple_chunks(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    monkeypatch.setattr(settings, "document_chunk_size_chars", 12)
    monkeypatch.setattr(settings, "document_chunk_overlap_chars", 2)
    token = get_access_token()

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "policy.txt",
                b"alpha beta gamma delta epsilon",
                "text/plain",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "ready"
    assert body["chunk_count"] > 1

    chunks_response = client.get(
        f"/api/v1/documents/{body['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert chunks_response.status_code == 200
    chunks = chunks_response.json()
    assert len(chunks) == body["chunk_count"]
    assert [chunk["chunk_index"] for chunk in chunks] == list(range(body["chunk_count"]))


def test_upload_markdown_document_extracts_heading_source(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "runbook.md",
                b"# Backup Runbook\n\nBackups run every night.\nRestore checks happen weekly.",
                "text/markdown",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "ready"
    assert body["chunk_count"] == 1

    chunks_response = client.get(
        f"/api/v1/documents/{body['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )
    chunks = chunks_response.json()

    assert chunks_response.status_code == 200
    assert chunks[0]["source_label"] == "Backup Runbook"
    assert chunks[0]["source_page"] is None

    search_response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "restore backup", "limit": 1},
    )

    assert search_response.status_code == 200
    assert search_response.json()[0]["source_label"] == "Backup Runbook"


def test_upload_docx_document_extracts_heading_source(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    docx_document = DocxDocument()
    docx_document.add_heading("Operations Handbook", level=1)
    docx_document.add_paragraph("Backups run every night.")
    docx_document.add_paragraph("Restore checks happen weekly.")
    buffer = BytesIO()
    docx_document.save(buffer)

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "handbook.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "ready"

    chunks_response = client.get(
        f"/api/v1/documents/{body['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )
    chunks = chunks_response.json()

    assert chunks_response.status_code == 200
    assert chunks[0]["source_label"] == "Operations Handbook"
    assert "Backups run every night" in chunks[0]["content"]


def test_upload_txt_document_can_enqueue_redis_ingestion(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    monkeypatch.setattr(settings, "document_ingestion_mode", "redis")
    queued_document_ids: list[int] = []

    def fake_enqueue_document_ingestion(redis_client, *, queue_name: str, document_id: int) -> None:
        assert queue_name == settings.document_ingestion_queue_name
        queued_document_ids.append(document_id)

    monkeypatch.setattr("app.api.v1.documents.get_redis_client", lambda: object())
    monkeypatch.setattr("app.api.v1.documents.enqueue_document_ingestion", fake_enqueue_document_ingestion)
    token = get_access_token()

    document = upload_txt_document(token)

    assert document["status"] == "pending"
    assert document["chunk_count"] == 0
    assert queued_document_ids == [document["id"]]

    chunks_response = client.get(
        f"/api/v1/documents/{document['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert chunks_response.status_code == 200
    assert chunks_response.json() == []


def test_worker_processes_queued_document_ingestion(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    monkeypatch.setattr(settings, "document_ingestion_mode", "redis")
    monkeypatch.setattr("app.api.v1.documents.get_redis_client", lambda: object())
    monkeypatch.setattr("app.api.v1.documents.enqueue_document_ingestion", lambda *args, **kwargs: None)
    token = get_access_token()
    document = upload_txt_document(token)

    with SessionLocal() as db:
        processed_document = process_document_ingestion(db, document["id"], settings=settings)

    assert processed_document is not None
    assert processed_document.status == "ready"
    assert processed_document.chunk_count == 1

    with SessionLocal() as db:
        duplicate_delivery = process_document_ingestion(db, document["id"], settings=settings)

    assert duplicate_delivery is not None
    assert duplicate_delivery.status == "ready"
    assert duplicate_delivery.chunk_count == 1

    search_response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "night backup", "limit": 5},
    )

    assert search_response.status_code == 200
    assert search_response.json()[0]["document_id"] == document["id"]


def test_owner_can_reindex_ready_document(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    document = upload_txt_document(token)

    response = client.post(
        f"/api/v1/documents/{document['id']}/reindex",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 202
    assert response.json()["status"] == "ready"
    assert response.json()["chunk_count"] == 1


def test_reindex_hides_other_users_document(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    owner_token = get_access_token()
    other_token = get_access_token()
    document = upload_txt_document(owner_token)

    response = client.post(
        f"/api/v1/documents/{document['id']}/reindex",
        headers={"Authorization": f"Bearer {other_token}"},
    )

    assert response.status_code == 404


def test_search_documents_returns_relevant_chunks(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    backup_document = upload_txt_document_with_content(
        token,
        b"Backups run every night. Restore checks happen weekly.",
        filename="backup-policy.txt",
    )
    upload_txt_document_with_content(
        token,
        b"Lunch menu includes pizza and salad.",
        filename="menu.txt",
    )

    response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "night backup restore", "limit": 2},
    )

    assert response.status_code == 200
    results = response.json()
    assert len(results) == 2
    assert results[0]["document_id"] == backup_document["id"]
    assert results[0]["document_filename"] == "backup-policy.txt"
    assert "Backups" in results[0]["content"]
    assert results[0]["score"] > results[1]["score"]
    with SessionLocal() as db:
        run = db.scalar(select(RetrievalRun).order_by(RetrievalRun.id.desc()))
    assert run is not None
    assert run.request_kind == "document_search"
    assert run.query_text is None
    assert run.candidate_count == 2
    assert run.selected_context_count == 0
    assert "content" not in run.candidates[0]
    assert run.model_versions == {"embedding": "fake-bow"}


def test_search_documents_requires_authentication() -> None:
    response = client.post(
        "/api/v1/documents/search",
        json={"query": "backup", "limit": 5},
    )

    assert response.status_code == 401


def test_reranked_search_persists_unavailable_backend_fallback(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    monkeypatch.setattr(settings, "reranker_backend", "disabled")
    token = get_access_token()
    upload_txt_document_with_content(token, b"Backups run every night.")

    response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "night backups", "limit": 1, "retrieval_strategy": "reranked"},
    )

    assert response.status_code == 200
    assert len(response.json()) == 1
    with SessionLocal() as db:
        run = db.scalar(select(RetrievalRun).order_by(RetrievalRun.id.desc()))
    assert run is not None
    assert run.strategy == "reranked"
    assert run.model_versions["reranker_outcome"] == "fallback"
    assert run.model_versions["reranker_candidate_count"] == 1
    assert run.model_versions["reranker_model"] == "bge-reranker-v2-m3"
    assert run.model_versions["reranker"].endswith(
        "@b5160aeac3c6c8fe7beaaaf04c9e0142826b58d1"
    )
    assert run.candidates[0]["strategy"] == "hybrid"
    assert "reranker" in run.timings_ms


def test_search_documents_hides_other_users_chunks(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    owner_token = get_access_token()
    other_token = get_access_token()
    upload_txt_document_with_content(owner_token, b"Secret project backup plan.")

    response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {other_token}"},
        json={"query": "secret backup", "limit": 5},
    )

    assert response.status_code == 200
    assert response.json() == []


def test_owner_can_grant_document_read_permission(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    owner_token = get_access_token()
    shared_email = f"shared-{uuid4().hex}@example.com"
    shared_token = create_user_and_get_token(shared_email)
    document = upload_txt_document_with_content(owner_token, b"Shared backup runbook.")

    grant_response = client.post(
        f"/api/v1/documents/{document['id']}/permissions",
        headers={"Authorization": f"Bearer {owner_token}"},
        json={"user_email": shared_email, "permission": "read"},
    )
    chunks_response = client.get(
        f"/api/v1/documents/{document['id']}/chunks",
        headers={"Authorization": f"Bearer {shared_token}"},
    )
    search_response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {shared_token}"},
        json={"query": "backup runbook", "limit": 5},
    )

    assert grant_response.status_code == 201
    assert grant_response.json()["permission"] == "read"
    assert chunks_response.status_code == 200
    assert chunks_response.json()[0]["content"] == "Shared backup runbook."
    assert search_response.status_code == 200
    assert search_response.json()[0]["document_id"] == document["id"]


def test_shared_document_permission_does_not_allow_delete(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    owner_token = get_access_token()
    shared_email = f"shared-{uuid4().hex}@example.com"
    shared_token = create_user_and_get_token(shared_email)
    document = upload_txt_document(owner_token)
    client.post(
        f"/api/v1/documents/{document['id']}/permissions",
        headers={"Authorization": f"Bearer {owner_token}"},
        json={"user_email": shared_email, "permission": "read"},
    )

    response = client.delete(
        f"/api/v1/documents/{document['id']}",
        headers={"Authorization": f"Bearer {shared_token}"},
    )

    assert response.status_code == 404


def test_search_documents_ignores_embeddings_from_another_model(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    upload_txt_document_with_content(token, b"Backups run every night.")
    monkeypatch.setattr(settings, "embedding_model", "replacement-model")

    response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "backups", "limit": 5},
    )

    assert response.status_code == 200
    assert response.json() == []


def test_reembed_all_document_chunks_replaces_embedding_model(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    document = upload_txt_document_with_content(token, b"Backups run every night.")

    with SessionLocal() as db:
        processed = reembed_all_document_chunks(
            db,
            provider=FakeEmbeddingProvider(model="replacement-model", dimensions=768),
            batch_size=1,
        )

    chunks_response = client.get(
        f"/api/v1/documents/{document['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert processed == 1
    assert chunks_response.status_code == 200
    assert chunks_response.json()[0]["embedding_model"] == "replacement-model"


def test_reembed_stale_only_skips_current_embeddings(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    upload_txt_document(token)

    with SessionLocal() as db:
        processed = reembed_all_document_chunks(
            db,
            provider=FakeEmbeddingProvider(model="fake-bow", dimensions=768),
            batch_size=1,
            stale_only=True,
        )

    assert processed == 0


def test_reembed_stale_only_replaces_mismatched_embedding_model(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    document = upload_txt_document(token)

    with SessionLocal() as db:
        chunk = db.scalars(select(DocumentChunk)).one()
        chunk.embedding_model = "old-model"
        db.commit()
        processed = reembed_all_document_chunks(
            db,
            provider=FakeEmbeddingProvider(model="fake-bow", dimensions=768),
            batch_size=1,
            stale_only=True,
        )

    chunks_response = client.get(
        f"/api/v1/documents/{document['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert processed == 1
    assert chunks_response.json()[0]["embedding_model"] == "fake-bow"


def test_reembed_rejects_wrong_embedding_dimensions(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    upload_txt_document(token)

    with SessionLocal() as db, pytest.raises(EmbeddingError, match="32 dimensions; expected 768"):
        reembed_all_document_chunks(
            db,
            provider=FakeEmbeddingProvider(model="wrong-size", dimensions=32),
            batch_size=1,
        )


def test_pgvector_schema_has_hnsw_cosine_index() -> None:
    with SessionLocal() as db:
        column_type = db.scalar(
            text(
                """
                SELECT format_type(a.atttypid, a.atttypmod)
                FROM pg_attribute AS a
                JOIN pg_class AS c ON c.oid = a.attrelid
                WHERE c.relname = 'document_chunks'
                  AND a.attname = 'embedding'
                  AND a.attnum > 0
                  AND NOT a.attisdropped
                """
            )
        )
        index_definition = db.scalar(
            text(
                """
                SELECT indexdef
                FROM pg_indexes
                WHERE tablename = 'document_chunks'
                  AND indexname = 'ix_document_chunks_embedding_hnsw_cosine'
                """
            )
        )

    assert column_type == "vector(768)"
    assert index_definition is not None
    assert "USING hnsw" in index_definition
    assert "vector_cosine_ops" in index_definition


def test_upload_document_requires_authentication(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))

    response = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "policy.txt",
                b"Backups run every night.",
                "text/plain",
            )
        },
    )

    assert response.status_code == 401
    assert not list(tmp_path.iterdir())


def test_upload_rejects_unsupported_content_type(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "notes.json",
                b"{}",
                "application/json",
            )
        },
    )

    assert response.status_code == 415
    assert not list(tmp_path.iterdir())


def test_upload_rejects_files_over_size_limit(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    monkeypatch.setattr(settings, "max_upload_size_bytes", 4)
    token = get_access_token()

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "policy.txt",
                b"too large",
                "text/plain",
            )
        },
    )

    assert response.status_code == 413
    assert not list(tmp_path.iterdir())


def test_upload_marks_unextractable_pdf_failed(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()

    response = client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "scan.pdf",
                b"not a real pdf",
                "application/pdf",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "failed"
    assert body["chunk_count"] == 0
    assert body["ingestion_error"]

    chunks_response = client.get(
        f"/api/v1/documents/{body['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert chunks_response.status_code == 200
    assert chunks_response.json() == []


def test_delete_document_removes_metadata_and_file(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    token = get_access_token()
    document = upload_txt_document(token)
    client.post(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "policy.txt",
                b"Retention reports run monthly.",
                "text/plain",
            )
        },
    )

    stored_files = list(tmp_path.iterdir())
    assert len(stored_files) == 2
    assert all(stored_file.exists() for stored_file in stored_files)

    response = client.delete(
        f"/api/v1/documents/{document['id']}",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 204
    assert response.content == b""
    assert not any(stored_file.exists() for stored_file in stored_files)

    list_response = client.get(
        "/api/v1/documents",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert list_response.status_code == 200
    assert list_response.json() == []

    chunks_response = client.get(
        f"/api/v1/documents/{document['id']}/chunks",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert chunks_response.status_code == 404

    versions_response = client.get(
        f"/api/v1/documents/{document['id']}/versions",
        headers={"Authorization": f"Bearer {token}"},
    )
    search_response = client.post(
        "/api/v1/documents/search",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "retention reports", "limit": 5},
    )
    with SessionLocal() as db:
        remaining_rows = db.execute(
            text(
                "SELECT "
                "(SELECT count(*) FROM documents WHERE id = :document_id), "
                "(SELECT count(*) FROM document_versions WHERE document_id = :document_id), "
                "(SELECT count(*) FROM document_chunks WHERE document_id = :document_id), "
                "(SELECT count(*) FROM document_permissions WHERE document_id = :document_id)"
            ),
            {"document_id": document["id"]},
        ).one()

    assert versions_response.status_code == 404
    assert search_response.status_code == 200
    assert search_response.json() == []
    assert tuple(remaining_rows) == (0, 0, 0, 0)


def test_delete_document_requires_authentication() -> None:
    response = client.delete("/api/v1/documents/1")

    assert response.status_code == 401


def test_delete_document_hides_other_users_document(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    owner_token = get_access_token()
    other_token = get_access_token()
    document = upload_txt_document(owner_token)

    response = client.delete(
        f"/api/v1/documents/{document['id']}",
        headers={"Authorization": f"Bearer {other_token}"},
    )

    assert response.status_code == 404
    assert len(list(tmp_path.iterdir())) == 1


def test_list_document_chunks_hides_other_users_document(
    tmp_path: Path,
    monkeypatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path))
    owner_token = get_access_token()
    other_token = get_access_token()
    document = upload_txt_document(owner_token)

    response = client.get(
        f"/api/v1/documents/{document['id']}/chunks",
        headers={"Authorization": f"Bearer {other_token}"},
    )

    assert response.status_code == 404
